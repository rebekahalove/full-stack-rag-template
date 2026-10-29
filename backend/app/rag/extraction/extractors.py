import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

from app.rag.types import ExtractedBlock, ExtractedDocument

PDF_TYPES = {"application/pdf"}
DOCX_TYPES = {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
MARKDOWN_TYPES = {"text/markdown", "text/x-markdown", "text/plain"}


class UnsupportedDocumentError(ValueError):
    pass


@dataclass(frozen=True, slots=True)
class ExtractionConfig:
    pdf_repeated_edge_min_pages: int
    pdf_repeated_edge_min_occurrences: int
    pdf_heading_max_characters: int


def _normalize(text: str) -> str:
    text = text.replace("\u00a0", " ").replace("\u00ad", "")
    text = re.sub(r"(?<=\w)-\n(?=\w)", "", text)
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _title(filename: str) -> str:
    return Path(filename).stem.replace("_", " ").replace("-", " ").strip()


def _make_block(
    content: str,
    *,
    content_type: str = "text",
    heading_path: list[str] | None = None,
    page: int | None = None,
    offset: int = 0,
    metadata: dict[str, object] | None = None,
) -> ExtractedBlock:
    return ExtractedBlock(
        content=content,
        content_type=content_type,
        heading_path=list(heading_path or []),
        page_start=page,
        page_end=page,
        offset_start=offset,
        offset_end=offset + len(content),
        metadata=dict(metadata or {}),
    )


def _extract_pdf(
    path: Path,
    filename: str,
    config: ExtractionConfig,
) -> ExtractedDocument:
    reader = PdfReader(path)
    raw_pages = [_normalize(page.extract_text() or "") for page in reader.pages]
    edge_counts: Counter[str] = Counter()
    if len(raw_pages) >= config.pdf_repeated_edge_min_pages:
        for text in raw_pages:
            lines = [line.strip() for line in text.splitlines() if line.strip()]
            edge_counts.update(lines[:1] + lines[-1:])
    repeated_edges = {
        line
        for line, count in edge_counts.items()
        if count >= config.pdf_repeated_edge_min_occurrences
    }

    blocks: list[ExtractedBlock] = []
    offset = 0
    headings: list[str] = []
    for page_number, raw_text in enumerate(raw_pages, start=1):
        lines = [line for line in raw_text.splitlines() if line.strip() not in repeated_edges]
        text = _normalize("\n".join(lines))
        for paragraph in re.split(r"\n\s*\n", text):
            paragraph = paragraph.strip()
            if not paragraph:
                continue
            content_type = (
                "heading"
                if len(paragraph) < config.pdf_heading_max_characters and paragraph.isupper()
                else "text"
            )
            if content_type == "heading":
                headings = [paragraph]
            blocks.append(
                _make_block(
                    paragraph,
                    content_type=content_type,
                    heading_path=headings,
                    page=page_number,
                    offset=offset,
                )
            )
            offset += len(paragraph) + 2
    metadata = {
        "page_count": len(reader.pages),
        "pdf_metadata": {str(key): str(value) for key, value in (reader.metadata or {}).items()},
    }
    return ExtractedDocument(_title(filename), path, blocks, metadata)


def _extract_docx(path: Path, filename: str) -> ExtractedDocument:
    document = DocxDocument(path)
    blocks: list[ExtractedBlock] = []
    headings: list[str] = []
    offset = 0
    paragraph_index = 0
    table_index = 0
    for element in document.element.body.iterchildren():
        if isinstance(element, CT_P):
            paragraph = Paragraph(element, document)
            content = _normalize(paragraph.text)
            if not content:
                continue
            style_name = paragraph.style.name if paragraph.style else ""
            match = re.match(r"Heading\s+(\d+)", style_name, re.IGNORECASE)
            content_type = "text"
            if match:
                level = max(1, int(match.group(1)))
                headings[level - 1 :] = [content]
                content_type = "heading"
            blocks.append(
                _make_block(
                    content,
                    content_type=content_type,
                    heading_path=headings,
                    offset=offset,
                    metadata={"paragraph_index": paragraph_index, "style": style_name},
                )
            )
            paragraph_index += 1
        elif isinstance(element, CT_Tbl):
            table = Table(element, document)
            rows = [[_normalize(cell.text) for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            width = max(len(row) for row in rows)
            normalized_rows = [row + [""] * (width - len(row)) for row in rows]
            header = normalized_rows[0]
            markdown = [
                "| " + " | ".join(header) + " |",
                "| " + " | ".join(["---"] * width) + " |",
            ]
            markdown.extend("| " + " | ".join(row) + " |" for row in normalized_rows[1:])
            content = "\n".join(markdown)
            blocks.append(
                _make_block(
                    content,
                    content_type="table",
                    heading_path=headings,
                    offset=offset,
                    metadata={"table_index": table_index, "row_count": len(rows)},
                )
            )
            table_index += 1
        else:
            continue
        offset += len(content) + 2
    return ExtractedDocument(
        _title(filename),
        path,
        blocks,
        {"paragraph_count": paragraph_index, "table_count": table_index},
    )


def _extract_markdown(path: Path, filename: str) -> ExtractedDocument:
    text = path.read_text(encoding="utf-8")
    blocks: list[ExtractedBlock] = []
    headings: list[str] = []
    buffer: list[str] = []
    block_type = "text"
    block_start = 0
    in_code = False

    def flush(end_offset: int) -> None:
        nonlocal buffer, block_start, block_type
        raw_content = "\n".join(buffer)
        content = raw_content.strip() if block_type == "code" else _normalize(raw_content)
        if content:
            block = _make_block(
                content,
                content_type=block_type,
                heading_path=headings,
                offset=block_start,
            )
            block.offset_end = end_offset
            blocks.append(block)
        buffer = []
        block_type = "text"

    offset = 0
    for line in text.splitlines(keepends=True):
        stripped = line.strip()
        if stripped.startswith("```"):
            if not in_code:
                flush(offset)
                block_start = offset
                block_type = "code"
            buffer.append(line.rstrip("\n"))
            in_code = not in_code
            offset += len(line)
            if not in_code:
                flush(offset)
            continue
        if in_code:
            buffer.append(line.rstrip("\n"))
            offset += len(line)
            continue
        heading = re.match(r"^(#{1,6})\s+(.+?)\s*#*$", stripped)
        if heading:
            flush(offset)
            level = len(heading.group(1))
            value = heading.group(2).strip()
            headings[level - 1 :] = [value]
            blocks.append(
                _make_block(value, content_type="heading", heading_path=headings, offset=offset)
            )
            offset += len(line)
            block_start = offset
            continue
        if not stripped:
            flush(offset)
            offset += len(line)
            block_start = offset
            continue
        if "|" in stripped:
            detected_type = "table"
        elif re.match(r"^([-*+] |\d+\. )", stripped):
            detected_type = "list"
        else:
            detected_type = "text"
        if buffer and detected_type != block_type:
            flush(offset)
            block_start = offset
        if not buffer:
            block_start = offset
            block_type = detected_type
        buffer.append(line.rstrip("\n"))
        offset += len(line)
    flush(len(text))
    return ExtractedDocument(_title(filename), path, blocks, {"character_count": len(text)})


def extract_document(
    path: Path,
    media_type: str,
    original_filename: str,
    config: ExtractionConfig,
) -> ExtractedDocument:
    suffix = path.suffix.lower() or Path(original_filename).suffix.lower()
    if media_type in PDF_TYPES or suffix == ".pdf":
        return _extract_pdf(path, original_filename, config)
    if media_type in DOCX_TYPES or suffix == ".docx":
        return _extract_docx(path, original_filename)
    if media_type in MARKDOWN_TYPES or suffix in {".md", ".markdown"}:
        return _extract_markdown(path, original_filename)
    raise UnsupportedDocumentError(f"Unsupported document type: {media_type or suffix}")
