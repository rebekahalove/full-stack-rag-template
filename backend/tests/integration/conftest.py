from __future__ import annotations

import io
import os
from collections.abc import AsyncIterator
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from types import SimpleNamespace
from uuid import UUID, uuid4

import httpx
import pytest
from alembic import command
from alembic.config import Config
from docx import Document as DocxDocument
from sqlalchemy.engine import make_url
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.core.config import Settings, get_settings
from app.main import create_app
from app.rag.types import AnswerResult, Citation
from app.services.documents import DocumentService


@dataclass
class RecordedIngestion:
    calls: list[tuple[UUID, Path, str, str]]

    async def ingest(
        self,
        document_id: UUID,
        path: Path,
        media_type: str,
        original_filename: str,
    ) -> None:
        self.calls.append((document_id, path, media_type, original_filename))


class DeterministicQuestionAnswering:
    async def answer(self, document_id: UUID, question: str) -> AnswerResult:
        if "dental" in question.lower():
            return AnswerResult(
                answer="The document does not contain that information.",
                insufficient_evidence=True,
                sources=[],
            )
        return AnswerResult(
            answer="Coverage begins on the first day of the month following enrollment.",
            insufficient_evidence=False,
            sources=[
                Citation(
                    chunk_id=uuid4(),
                    document_id=document_id,
                    document_name="benefits.md",
                    excerpt="Coverage begins on the first day of the month following enrollment.",
                    heading_path=["Benefits handbook", "Eligibility"],
                    page_start=None,
                    page_end=None,
                    score=0.94,
                )
            ],
        )


class InMemoryDocumentRepository:
    def __init__(self) -> None:
        self.documents: dict[UUID, SimpleNamespace] = {}

    async def create(self, **values: object) -> SimpleNamespace:
        now = datetime.now(UTC)
        document = SimpleNamespace(
            id=uuid4(),
            chunk_count=2,
            error_message=None,
            created_at=now,
            updated_at=now,
            **values,
        )
        self.documents[document.id] = document
        return document

    async def list(self) -> list[SimpleNamespace]:
        return list(self.documents.values())

    async def get(self, document_id: UUID) -> SimpleNamespace | None:
        return self.documents.get(document_id)

    async def delete(self, document_id: UUID) -> bool:
        return self.documents.pop(document_id, None) is not None

    async def update_status(
        self,
        document_id: UUID,
        status: str,
        error_message: str | None = None,
    ) -> SimpleNamespace | None:
        document = self.documents.get(document_id)
        if document is not None:
            document.status = status
            document.error_message = error_message
            document.updated_at = datetime.now(UTC)
        return document

    async def count_chunks(self, document_id: UUID) -> int:
        document = self.documents.get(document_id)
        return int(document.chunk_count) if document else 0


class HealthyDatabase:
    async def check(self) -> dict[str, str]:
        return {"status": "ok", "database": "ok"}


@pytest.fixture
def ingestion() -> RecordedIngestion:
    return RecordedIngestion(calls=[])


@pytest.fixture
def repository() -> InMemoryDocumentRepository:
    return InMemoryDocumentRepository()


@pytest.fixture
async def api_client(
    tmp_path: Path,
    repository: InMemoryDocumentRepository,
    ingestion: RecordedIngestion,
) -> AsyncIterator[httpx.AsyncClient]:
    settings = Settings(
        upload_directory=tmp_path / "uploads",
        openai_api_key=None,
        max_upload_size_mb=0.1,
    )
    service = DocumentService(
        repository,
        ingestion,
        DeterministicQuestionAnswering(),
        upload_directory=settings.upload_directory,
        max_upload_size_mb=settings.max_upload_size_mb,
        upload_read_size_bytes=settings.upload_read_size_bytes,
        upload_pdf_header_bytes=settings.upload_pdf_header_bytes,
        allowed_extensions=settings.upload_allowed_extensions,
        question_max_characters=settings.question_max_characters,
    )
    app = create_app(
        settings=settings,
        document_service=service,
        health_checker=HealthyDatabase(),
    )
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
        yield client


def minimal_docx_bytes(text: str = "Coverage begins after enrollment.") -> bytes:
    target = io.BytesIO()
    document = DocxDocument()
    document.add_heading("Benefits", level=1)
    document.add_paragraph(text)
    document.save(target)
    return target.getvalue()


def minimal_pdf_bytes(text: str = "Coverage begins after enrollment") -> bytes:
    """Build a one-page PDF with a text stream and a correct cross-reference table."""
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"
        ),
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = io.BytesIO()
    output.write(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(output.tell())
        output.write(f"{index} 0 obj\n".encode())
        output.write(obj)
        output.write(b"\nendobj\n")
    xref = output.tell()
    output.write(f"xref\n0 {len(objects) + 1}\n".encode())
    output.write(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.write(f"{offset:010d} 00000 n \n".encode())
    output.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return output.getvalue()


@pytest.fixture(scope="session")
def migrated_database_url() -> str:
    """Apply migrations only when an explicit disposable integration URL is supplied."""
    configured = os.getenv("TEST_DATABASE_URL")
    if not configured:
        pytest.skip("TEST_DATABASE_URL is not set; skipping PostgreSQL integration tests")

    url = make_url(configured)
    if url.get_backend_name() != "postgresql":
        pytest.fail("TEST_DATABASE_URL must point to PostgreSQL with pgvector")
    database_name = url.database or ""
    if "test" not in database_name.casefold():
        pytest.fail(
            "TEST_DATABASE_URL database name must contain 'test' to protect development data"
        )
    if url.drivername != "postgresql+asyncpg":
        url = url.set(drivername="postgresql+asyncpg")
    async_url = url.render_as_string(hide_password=False)

    backend_root = Path(__file__).parents[2]
    migrations = backend_root / "migrations"
    if not migrations.exists():
        pytest.fail("Alembic migrations are missing")
    config_path = backend_root / "alembic.ini"
    config = Config(str(config_path)) if config_path.exists() else Config()
    config.set_main_option("script_location", str(migrations))
    config.set_main_option("sqlalchemy.url", async_url.replace("%", "%%"))
    previous_database_url = os.environ.get("DATABASE_URL")
    os.environ["DATABASE_URL"] = async_url
    get_settings.cache_clear()
    try:
        command.upgrade(config, "head")
    finally:
        if previous_database_url is None:
            os.environ.pop("DATABASE_URL", None)
        else:
            os.environ["DATABASE_URL"] = previous_database_url
        get_settings.cache_clear()

    return async_url


@pytest.fixture
async def db_session(migrated_database_url: str) -> AsyncIterator[AsyncSession]:
    engine = create_async_engine(migrated_database_url, pool_pre_ping=True)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    async with factory() as session:
        yield session
        await session.rollback()
    await engine.dispose()
