from pathlib import Path

from docx import Document as DocxDocument

from app.rag.extraction import ExtractionConfig, extract_document

EXTRACTION_CONFIG = ExtractionConfig(
    pdf_repeated_edge_min_pages=3,
    pdf_repeated_edge_min_occurrences=3,
    pdf_heading_max_characters=100,
)


def test_markdown_extraction_preserves_structure_and_offsets(tmp_path: Path) -> None:
    source = tmp_path / "guide.md"
    source.write_text(
        "# Enrollment\n\nChoose a plan during open enrollment.\n\n"
        "## Dates\n\n- Opens October 1\n- Closes October 31\n",
        encoding="utf-8",
    )

    document = extract_document(source, "text/markdown", "guide.md", EXTRACTION_CONFIG)

    assert document.title == "guide"
    assert [block.content_type for block in document.blocks] == [
        "heading",
        "text",
        "heading",
        "list",
    ]
    assert document.blocks[-1].heading_path == ["Enrollment", "Dates"]
    assert all(block.offset_start is not None for block in document.blocks)


def test_markdown_extraction_keeps_fenced_code_together(tmp_path: Path) -> None:
    source = tmp_path / "example.md"
    source.write_text("# Example\n\n```python\nprint('hello')\n```\n", encoding="utf-8")

    document = extract_document(source, "text/markdown", "example.md", EXTRACTION_CONFIG)

    code = next(block for block in document.blocks if block.content_type == "code")
    assert "print('hello')" in code.content


def test_docx_extraction_keeps_table_order_headers_and_section(tmp_path: Path) -> None:
    source = tmp_path / "plans.docx"
    document = DocxDocument()
    document.add_heading("Plan costs", level=1)
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Plan"
    table.rows[0].cells[1].text = "Monthly cost"
    table.rows[1].cells[0].text = "Standard"
    table.rows[1].cells[1].text = "$50"
    document.add_paragraph("Costs are deducted from payroll.")
    document.save(source)

    extracted = extract_document(
        source,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        source.name,
        EXTRACTION_CONFIG,
    )

    assert [block.content_type for block in extracted.blocks] == ["heading", "table", "text"]
    assert extracted.blocks[1].heading_path == ["Plan costs"]
    assert "| Plan | Monthly cost |" in extracted.blocks[1].content
    assert extracted.metadata["table_count"] == 1
